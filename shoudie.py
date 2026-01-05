from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def add_slide(prs, title_str, content_str):
    slide_layout = prs.slide_layouts[1]  # 使用标题和内容布局
    slide = prs.slides.add_slide(slide_layout)
    
    # 设置背景颜色为深蓝色 (手碟的深邃感)
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(10, 17, 40) 

    # 设置标题样式
    title = slide.shapes.title
    title.text = title_str
    title_text_frame = title.text_frame.paragraphs[0]
    title_text_frame.font.color.rgb = RGBColor(225, 206, 122) # 金色
    title_text_frame.font.bold = True

    # 设置正文样式
    content = slide.placeholders[1]
    content.text = content_str
    for paragraph in content.text_frame.paragraphs:
        paragraph.font.color.rgb = RGBColor(255, 255, 255) # 白色文字
        paragraph.font.size = Pt(18)

# 创建 PPT 实例
prs = Presentation()

# 幻灯片 1：封面
slide_layout = prs.slide_layouts[0]
slide = prs.slides.add_slide(slide_layout)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = RGBColor(10, 17, 40)
title = slide.shapes.title
title.text = "手碟入门第一课：初次对话"
subtitle = slide.placeholders[1]
subtitle.text = "认识你的声音伙伴——D Kurd调式手碟\n讲师：[您的姓名]"

# 幻灯片 2：大纲
add_slide(prs, "今天我们将一起经历…", 
          "1. 开启旅程 (15min) - 手碟的故事\n"
          "2. 初次触碰 (25min) - 姿势与手法\n"
          "3. 声音地图 (35min) - D Kurd 音阶探索\n"
          "4. 总结与启程 (15min) - 课后练习")

# 幻灯片 3：第一部分
add_slide(prs, "认识你的新朋友", 
          "• 出生：21世纪初，瑞士\n"
          "• 灵感：钢鼓、陶罐鼓等世界乐器\n"
          "• 使命：创造一种全新的、亲密的、易于冥想的声音\n"
          "• D Kurd 调式：深邃、宁静、略带神秘")

# ... 可以按照这个模式继续添加剩下的幻灯片 ...
# 考虑到代码长度，这里省略中间重复部分，逻辑一致

# 保存文件
prs.save('手碟第一课_课件.pptx')
print("PPT 已生成：手碟第一课_课件.pptx")
